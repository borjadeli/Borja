"""Genera un .docx corporativo (Banking & Finance) que RESPONDE a las dos consultas
del equipo Quant (Borja) sobre BORJA:

  (1) Lógica de selección + funnel (cuántas empresas pasan cada filtro, sectores,
      market cap) y "qué elige la IA al armar el portfolio".
  (2) Precio de referencia (análisis a precio actual) + propuesta de recambio
      ("si encuentra algo mejor, propone el cambio").

Incluye el análisis del equipo y deja explícitas las DOS decisiones que corresponden
al equipo Quant. Auto-contenido para compartir.

Reutiliza el toolkit de estilo de ops/generar_doc_evaluacion_quant.py (misma paleta).

Output: data/reportes/borja_respuesta_consultas_<timestamp>.docx
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

# Reutilizamos la paleta y los helpers del generador de evaluación (mismo look & feel).
from ops.generar_doc_evaluacion_quant import (
    NAVY, ACCENT, SLATE_900, SLATE_700, SLATE_500, EMERALD, CRIMSON,
    add_table, add_h1, add_h2, add_h3, add_p, add_bullet, add_callout, add_footer,
)

_MESES_ES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
             "agosto", "septiembre", "octubre", "noviembre", "diciembre")


def _fecha_es() -> str:
    hoy = datetime.now()
    return f"{hoy.day:02d} de {_MESES_ES[hoy.month - 1]} de {hoy.year}"


def _quote(doc, text) -> None:
    """Cita textual de la consulta de Borja, en recuadro tenue con acento gris."""
    add_callout(doc, None, [text], fill="F8FAFC", accent_hex="64748B")


def build() -> Document:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    add_footer(doc)

    # ============================ CARÁTULA ============================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run("BORJA")
    r.bold = True
    r.font.size = Pt(46)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Núcleo Determinista · v3")
    r.italic = True
    r.font.size = Pt(13)
    r.font.color.rgb = SLATE_500

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    r = p.add_run("Lógica de Selección, Funnel\ny Precio de Referencia")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = SLATE_900

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Respuesta a las consultas del equipo Quant + dos decisiones a acordar")
    r.font.size = Pt(11)
    r.font.color.rgb = SLATE_500

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(150)
    for label, val in (("Destinatario", "Borja — Equipo de Quants"),
                       ("Fecha", _fecha_es()),
                       ("Clasificación", "Confidencial — Uso Interno")):
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = pp.add_run(f"{label}:  ")
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = NAVY
        r2 = pp.add_run(val)
        r2.font.size = Pt(10)
        r2.font.color.rgb = SLATE_700

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("Elaborado y verificado contra el estado actual del código el 02-jul-2026 "
                  "(paquetes backend/nucleo y backend/seleccion) y el diseño canónico (docs 11_* en "
                  "init_docs). Complementa la «Revisión Cuantitativa de Reglas de Entrada y Salida» del 30-jun.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = SLATE_500
    doc.add_page_break()

    # ==================== 1. RESUMEN EJECUTIVO ====================
    add_h1(doc, "Resumen Ejecutivo", kicker="Sección 1")
    add_callout(doc, "Veredicto en una línea", [
        "Las dos observaciones son correctas y, en su mayoría, ya están contempladas por el "
        "rediseño determinista que estamos migrando. Hay un reencuadre importante en el punto (1) "
        "y una decisión de diseño genuina en el punto (2).",
        "Punto (1) — «qué elige la IA»:  en el núcleo nuevo la IA NO elige la cartera. La elección "
        "es 100% determinista y el LLM solo aporta hechos fechados con fuente. Tu pedido, entonces, "
        "no es «ponerle criterio a la IA» (el criterio ya es nuestro, son reglas) sino tener "
        "VISIBILIDAD de ese criterio. Eso es un informe de transparencia del funnel, y es "
        "perfectamente escribible: el dato ya se genera etapa por etapa.",
        "Punto (2) — «que proponga el cambio si hay algo mejor»:  tu intuición sobre el precio "
        "actual es exacta — es la primera invariante del sistema. Y el «proponé el recambio» es "
        "medio-cierto hoy: el ranking se recomputa cada ciclo, pero la regla explícita que desaloja "
        "a la posición más débil cuando aparece una candidata superior está deliberadamente apagada. "
        "Reactivarla —o no— es la decisión de diseño a acordar.",
    ])
    add_p(doc, "El resto del documento desarrolla cada punto, describe el informe propuesto y "
               "cierra con las dos decisiones que corresponden al equipo Quant.", color=SLATE_500,
          italic=True)

    # ==================== 2. PUNTO 1 ====================
    add_h1(doc, "Punto 1 — Qué elige el sistema al armar la cartera", kicker="Consulta de Borja")
    _quote(doc, "«Necesito saber bien la lógica de cómo está laburando. Primero cuántas empresas "
                "llegan al primer filtro, de qué sectores, qué market cap tienen. Lo mismo con las "
                "que superan el segundo filtro, tener un buen informe del detalle de esas empresas. "
                "Y por último y clave, entender cómo y qué elige la IA al seleccionar el portfolio. "
                "Eso es clave porque lo dejé a criterio de la IA y ahí es donde capaz le tenemos que "
                "poner criterio nuestro también.»")

    add_h2(doc, "2.1 · El reencuadre clave: en el núcleo nuevo la IA no elige")
    add_p(doc, "El objetivo central del rediseño fue justamente sacar al LLM de la decisión. Hoy el "
               "LLM es un traedor de hechos fechados con fuente: aporta datos —incluidos los precios "
               "bull/bear y las señales de calidad, que son el input más blando (así lo marca la "
               "evaluación del 30-jun)—, pero nunca asigna arquetipo, tier, mérito ni decisión. La selección "
               "de la cartera es una cadena de tablas deterministas. Por lo tanto, la parte de tu "
               "consulta — «lo dejé a criterio de la IA y ahí capaz le tenemos que poner criterio "
               "nuestro» — ya está resuelta a nivel de diseño: el criterio ya es nuestro. Lo que "
               "falta no es gobernar a la IA, es ver el criterio en acción.")
    add_callout(doc, None, [
        "Esto reencuadra tu pedido: de «controlar a la IA» a «tunear nuestros parámetros "
        "explícitos» (cortes de Quality, umbrales de asimetría, Hard Caps, caps de concentración). "
        "Que es exactamente lo que un quant quiere poder hacer — y lo que el informe de funnel "
        "te va a habilitar.",
    ])

    add_h2(doc, "2.2 · La cadena de decisión (determinista, de punta a punta)")
    add_p(doc, "La cartera se arma recorriendo esta secuencia. Cada etapa es inspeccionable y es "
               "donde vive uno de tus «filtros»:")
    add_table(doc,
        ["Etapa", "Qué hace", "Tu «filtro»"],
        [["Universo", "Small/mid-cap de industrias del futuro; screen de liquidez > $1M/día.", "Población inicial"],
         ["Filtro 1 · Arquetipo + elegibilidad", "Clasifica cada empresa en uno de 5 arquetipos (Madura / Growth / Tracción / Pre-revenue / Binary) y ajusta por concentración de cliente (ex-top). El descarte por evidencia se hace en el Filtro 2.", "1.er filtro"],
         ["Filtro 2 · Gate de evidencia", "Checklist pasa/no-pasa duro, sin compensación (todos los ítems deben pasar): triggers con fuente, contrato/design-in económico y —según arquetipo— financiación-hasta-el-hito. Los roll-ups (crecimiento por M&A) caen acá por G1.", "2.º filtro"],
         ["Ranking · Mérito", "Mérito = Quality (tier ordinal 1-4) × Asimetría al precio actual. Cada empresa recibe además un Hard Cap = tope por arquetipo × Fragility (severidad del downside), topeado por CRR.", "Selección"],
         ["Selección · Selector + Sizing", "Selector: top-N por mérito (normalización relativa). Sizing: acota al Hard Cap, aplica piso 5% y el cubo de fragilidad. Máx 12 posiciones.", "Cartera final"]],
        col_widths_cm=[4.4, 9.6, 2.5])

    add_h2(doc, "2.3 · El funnel que pedís — mapeado a lo que el sistema ya produce")
    add_p(doc, "Cada cosa que pedís corresponde a una etapa y ya se captura en cada corrida:")
    add_table(doc,
        ["Lo que pedís", "Etapa del núcleo", "Qué expone el informe"],
        [["Cuántas al 1.er filtro; de qué sectores; qué market cap",
          "Universo → Arquetipo + elegibilidad",
          "Conteo total; breakdown por sector y por bucket de market cap; y qué queda fuera de mandato hoy (destapa los mega-caps, que hoy no tienen techo)."],
         ["Cuántas superan el 2.º filtro + detalle de esas empresas",
          "Gate de evidencia (U1/U2 + M/G/T/P/B + financiación)",
          "Pasa/no-pasa por empresa, con el ítem exacto que falló. Éste es el «detalle» que pedís."],
         ["Cómo y qué elige al armar el portfolio",
          "Ranking de mérito + Selector",
          "Tabla ordenada por mérito con sizing: muestra por qué entró A y no B, con números."]],
        col_widths_cm=[5.0, 4.6, 6.9])
    add_p(doc, "En la última corrida real sobre 37 tickers, el funnel arrojó 29 candidatas / 8 "
               "rechazos, con el dato por etapa ya guardado en disco. Falta la capa de presentación, "
               "no el dato: es un trabajo de reporte/BI, no de algoritmo.", color=SLATE_900)

    add_h2(doc, "2.4 · Ojo: hoy conviven dos sistemas")
    add_callout(doc, "Sobre cuál sistema hacemos el informe", [
        "El sistema VIEJO (score 0-10, donde el LLM sí puntúa — probablemente lo que estás mirando) "
        "sigue corriendo. El núcleo NUEVO determinista corre en paper y lo reemplaza cuando valide "
        "(gate PASS_RATE 80%).",
        "Antes de escribir el informe hay que decidir sobre cuál lo hacemos. Recomendación: el "
        "nuevo — es el que vamos a operar y donde el criterio ya es nuestro. Ver Decisión A (§5).",
    ], fill="EFF6FF", accent_hex="1E3A8A")

    add_h2(doc, "2.5 · Un hallazgo que el informe destapa solo")
    add_p(doc, "Tu pregunta «qué market cap tienen» no es cosmética. El universo hoy no aplica techo "
               "de market-cap ni filtro temático, y por eso se colaron mega-caps fuera de mandato: IBM "
               "está en el universo cacheado y AAPL incluso entró al libro (12.ª posición, arquetipo "
               "Madura). El informe de funnel lo expone en la primera sección: además de dar "
               "transparencia, sirve de control de calidad y justifica la higiene de universo pendiente "
               "(techo de market-cap + filtro temático).")

    # ==================== 3. PUNTO 2 ====================
    add_h1(doc, "Punto 2 — Precio de referencia y propuesta de recambio", kicker="Consulta de Borja")
    _quote(doc, "«La entrada y la salida las podemos flexibilizar. Lo más importante es que, si "
                "encuentra una empresa mejor que alguna del portfolio, te propone el cambio. Si no te "
                "propone cambio, es que no tiene nada mejor. Y las empresas las analiza con precio "
                "actual: si el precio subió 50% y no la reemplaza, es que cree que aun con esa subida "
                "hay valor.»")

    add_h2(doc, "3.1 · El precio actual — tu intuición es exacta (y ya es la primera invariante)")
    add_p(doc, "Clavaste cómo funciona. La asimetría se recalcula al precio vigente antes de entrar "
               "o aumentar; es la Invariante #1 del sistema. Que una acción suba 50% y no la reemplace "
               "significa, literalmente, que el núcleo recomputó el upside al precio nuevo y todavía "
               "encontró margen sobre el umbral. No es inercia: es el mecanismo. Y tiene una propiedad "
               "linda: una posición que cae con la tesis intacta ve subir su asimetría, lo que "
               "auto-corrige el sesgo de comprar caro / vender barato.")

    add_h2(doc, "3.2 · «Proponer el cambio» — hoy es medio-cierto")
    add_h3(doc, "Lo que ya funciona", color=EMERALD)
    add_bullet(doc, "En cada ciclo el núcleo recalcula el mérito de todo el universo al precio "
                    "actual, y el Selector ordena a las candidatas por mérito.")
    add_bullet(doc, "Existen 5 vías de salida. Están vivas las de precio —vía 4 (asimetría muerta) "
                    "y vía 5 (tamaño)— y, ya cableadas de forma determinista, la vía 2 (CRR-crítico) "
                    "y la vía 3 (red flags blandos). Cuando una posición se rompe, sale, y su lugar "
                    "queda disponible para la mejor candidata.")
    add_bullet(doc, "Cuando hay un slot libre, la mejor candidata por mérito entra.")
    add_h3(doc, "Lo que hoy NO ocurre", color=CRIMSON)
    add_bullet(doc, "No hay una regla explícita «challenger» que desaloje a la posición más débil "
                    "cuando aparece una candidata de mayor mérito y la cartera está llena.")
    add_bullet(doc, "El rediseño eliminó a propósito las 8 reglas de rotación del sistema viejo; el "
                    "default quedó en «mantener salvo que se rompa una vía».")
    add_bullet(doc, "La config actual es forward-only + posiciones grandfathered + piramidación "
                    "apagada (flag PERMITIR_PIRAMIDAR = False). La cartera está hoy 12/12 y 100% "
                    "desplegada; como las entradas solo ocupan slots libres, no queda margen para que "
                    "una candidata mejor entre: hasta que un titular salga por una vía de salida, la "
                    "cartera es estática.")
    add_callout(doc, None, [
        "En criollo: hoy el sistema no te va a proponer «sacá X y meté Y porque Y es mejor». Va a "
        "sacar a X solo si X se rompe: por asimetría (vía 4) o tamaño (vía 5), y —ya cableadas de "
        "forma determinista— por CRR-crítico (vía 2) o red flags blandos ≥3 (vía 3). Todavía faltan "
        "la invalidación / stop de progreso (vía 1) y la governance hard, que necesitan el pipeline "
        "de monitor/noticias. Recién cuando un titular sale por una de esas vías entra la mejor "
        "candidata. El recambio por mérito que describís es un diseño que podemos activar, no un "
        "bug: es la Decisión B (§5).",
    ], fill="FFFBEB", accent_hex="D97706")

    add_h2(doc, "3.3 · Si lo activamos: challenger con histéresis (recomendación quant)")
    add_p(doc, "Un «swap si mérito_nuevo > mérito_viejo» ingenuo genera churn que se come el alpha "
               "(costos de transacción, impacto de liquidez en small-caps, ventana de protección de "
               "timing). Recomendación: que la candidata tenga que superar a la titular por un MARGEN "
               "—una banda de histéresis, no por un pelo— y respetando la protección de timing 45-60d. "
               "Dato clave: esta lógica de recambio por mérito YA EXISTE en el código, pero aplicada a "
               "la watchlist (watchlist.py: «desplaza a la peor no-protegida por mérito solo si la nueva "
               "la supera»), no a la cartera. Activar la Decisión B es, en buena medida, extender ese "
               "mismo patrón al libro. Además el sistema ya usa mecanismos de histéresis / bloqueo "
               "reutilizables: la banda de tolerancia 1.25× de la vía 5, los candados A (bloquea "
               "re-gate) y B (re-entrada de una sola vida), y el umbral 30%/20% del monitor.")
    add_p(doc, "Nota de equipo: esto roza una preferencia durable — «disciplina sobre actividad, "
               "mejor un slot vacío que forzar una entrada». El challenger sube el turnover. Por eso "
               "es una decisión de equipo, no técnica.", italic=True, color=SLATE_500)

    # ==================== 4. EL INFORME ====================
    add_h1(doc, "El Informe de Transparencia del Funnel", kicker="Sección 4")
    add_p(doc, "Sí, se puede escribir — y es acotado. Propongo esta estructura, por corrida:")
    add_table(doc,
        ["Sección", "Contenido", "Responde a"],
        [["1 · Universo", "Conteo + breakdown por sector y por bucket de market cap.", "«cuántas / qué sectores / qué market cap» (destapa mega-caps)"],
         ["2 · Filtro 1 (Arquetipo/elegibilidad)", "Cuántas pasan + cada descarte con su motivo.", "«cuántas llegan al primer filtro»"],
         ["3 · Filtro 2 (Gate de evidencia)", "Cuántas pasan + tabla por empresa con el ítem que falló.", "«detalle de las que superan el segundo filtro»"],
         ["4 · Ranking de mérito", "Tabla ordenada (Quality × asimetría al precio actual) + sizing + cubo.", "«cómo y qué elige»"],
         ["5 · Cartera + concentración + challenger", "Qué candidata supera a qué titular y por cuánto margen.", "Punto 2 (visual)"]],
        col_widths_cm=[4.6, 7.4, 4.5])
    add_p(doc, "La mayoría del dato crudo ya se genera en cada corrida (shadow del núcleo + libro "
               "paper y su journal). Falta la capa de presentación. Formato posible: página nueva en "
               "el dashboard (estilo Banking/Finance, p. ej. /transparencia) o un DOCX/PDF generado — "
               "el mismo pipeline que produjo este documento sirve de base.")

    # ==================== 5. DECISIONES ====================
    add_h1(doc, "Dos Decisiones del Equipo Quant", kicker="Sección 5")
    add_p(doc, "Ninguna es técnica; las dos definen qué se construye y cómo opera el sistema. "
               "Corresponden al equipo (Borja + Quants).")

    add_callout(doc, "Decisión A — ¿Sobre qué sistema hacemos el informe?", [
        "Opción A1 (recomendada):  el núcleo NUEVO determinista — es el que vamos a operar y donde "
        "el criterio ya es nuestro. Bonus: el informe sirve también como validación previa al corte "
        "(PASS_RATE).",
        "Opción A2:  el sistema VIEJO (score 0-10 con LLM) — es lo que corre en producción hoy, pero "
        "está en camino de ser reemplazado.",
    ], fill="EFF6FF", accent_hex="1E3A8A")

    add_callout(doc, "Decisión B — ¿Activamos la regla «challenger» (recambio por mérito)?", [
        "Opción B1:  activar el challenger CON histéresis — la cartera propone recambios cuando una "
        "candidata supera a una titular por un margen. Más cercano a lo que describe Borja; sube el "
        "turnover. (En buena medida es extender a la cartera el recambio por mérito que la watchlist "
        "ya aplica.)",
        "Opción B2 (statu quo):  mantener «hold hasta que se rompa una vía» — más disciplinado, menos "
        "churn, alineado con «mejor un slot vacío que forzar».",
        "Trade-off:  B1 captura el upside del recambio pero paga costos de transacción y riesgo de "
        "churn en small-caps; B2 es más pasivo, pero puede quedarse con una titular mediocre mientras "
        "hay algo mejor esperando.",
    ], fill="FFFBEB", accent_hex="D97706")

    # ==================== 6. PRÓXIMOS PASOS ====================
    add_h1(doc, "Próximos Pasos", kicker="Cierre")
    add_bullet(doc, "Resolver la Decisión A y la Decisión B.", bold_lead="Equipo:  ")
    add_bullet(doc, "Con A resuelta, armar el informe de funnel sobre una corrida real y compartirlo.",
               bold_lead="Nosotros:  ")
    add_bullet(doc, "Con B resuelta: si es B1, cablear el challenger con histéresis (+ tests); si es "
                    "B2, dejarlo documentado como regla explícita.", bold_lead="Nosotros:  ")
    add_bullet(doc, "Higiene de universo (techo de market-cap + filtro temático) para sacar los "
                    "mega-caps off-mandate. Ya identificado, va en paralelo.", bold_lead="Paralelo:  ")
    add_p(doc, "Cualquiera de las dos decisiones se puede cambiar después; ninguna es irreversible. "
               "Lo importante es fijar el punto de partida.", italic=True, color=SLATE_900)

    return doc


def main() -> Path:
    out_dir = Path(__file__).resolve().parent.parent / "data" / "reportes"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"borja_respuesta_consultas_{datetime.now().strftime('%Y%m%dT%H%M%S')}.docx"
    build().save(out)
    return out


if __name__ == "__main__":
    path = main()
    print(f"OK -> {path}")
